"""Documents commerciaux : devis et facture (HTML modifiables + PDF exemples),
signature e-mail, papier à en-tête Word, carte des prestations A4, présentation
commerciale 16:9."""
import json
import subprocess
from html import escape
from pathlib import Path

from build_icons import LINEA, SERVICES
from kit import A, B, G, N, ROOT, TR, V, font_css, logo_svg
from textes import AVIS, METHODE, PITCH_COURT, PITCH_LONG, PRESTATIONS, SIGNATURE

OUT = ROOT / "08_papeterie"
DOCS = OUT / "documents"
FONT_REL = "../../03_typographie/polices"
IMG_REL = "../../06_photographie/rendus-3d/"


def icon(name, color=N, size=28):
    body, title = SERVICES[name]
    return (f'<svg class="ico" viewBox="0 0 24 24" width="{size}" height="{size}" fill="none" stroke="{color}" '
            f'stroke-width="1.5" stroke-linecap="square" aria-hidden="true">{body}{LINEA}</svg>')


BASE_CSS = f"""
:root{{--nero:{N};--bianco:{B};--verde:{V};--verdaccio:{G};--argento:{A};--travertino:{TR}}}
*{{box-sizing:border-box}}
body{{margin:0;font-family:Archivo,Arial,sans-serif;color:var(--nero);-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.eyebrow{{font-family:Archivo,Arial,sans-serif;font-stretch:125%;font-weight:600;text-transform:uppercase;letter-spacing:.22em;font-size:7pt}}
.mono{{font-family:'DM Mono',monospace;letter-spacing:.04em}}
.serif{{font-family:'Bodoni Moda',Didot,serif;font-weight:400}}
.linea{{height:0;border-top:.6pt solid var(--nero)}}
"""

# ---------------------------------------------------------------- devis / facture
LINES = [
    ("Lavage de préparation et décontamination (ferreuse, goudron, argile)", 1, 149.00),
    ("Correction de vernis — 2 étapes, contrôle sous lumière", 1, 690.00),
    ("Protection céramique carrosserie — formule 5 ans", 1, 890.00),
    ("Jantes, étriers et vitrage — protection hydrophobe", 1, 120.00),
]


def money(v):
    return f"{v:,.2f}".replace(",", " ").replace(".", ",") + " €"


def invoice(kind="facture"):
    is_q = kind == "devis"
    title = "Devis" if is_q else "Facture"
    num = "D-2026-0147" if is_q else "F-2026-0147"
    rows = "".join(
        f'<tr><td contenteditable="true">{escape(d)}</td><td class="num qty" contenteditable="true">{q}</td>'
        f'<td class="num pu" contenteditable="true">{money(p)}</td><td class="num">20 %</td>'
        f'<td class="num tot">{money(q * p)}</td></tr>' for d, q, p in LINES)
    ht = sum(q * p for _, q, p in LINES)
    extra_q = """
      <section class="cond">
        <div><p class="eyebrow">Validité</p><p contenteditable="true">30 jours à compter de la date d'émission.</p></div>
        <div><p class="eyebrow">Immobilisation estimée</p><p contenteditable="true">3 jours ouvrés, dépôt le matin.</p></div>
        <div><p class="eyebrow">Acompte</p><p contenteditable="true">30 % à la réservation du créneau.</p></div>
      </section>
      <section class="accord">
        <p class="eyebrow">Bon pour accord</p>
        <p>Date, signature et mention « Bon pour accord »</p>
        <div class="sig"></div>
      </section>""" if is_q else """
      <section class="cond">
        <div><p class="eyebrow">Règlement</p><p contenteditable="true">Réglée le 24.10.2026 par virement.</p></div>
        <div><p class="eyebrow">Échéance</p><p contenteditable="true">À réception.</p></div>
        <div><p class="eyebrow">Garantie céramique</p><p contenteditable="true">Sous réserve d'un entretien annuel à l'atelier.</p></div>
      </section>"""
    return f"""<!doctype html>
<html lang="fr"><head><meta charset="utf-8"><title>Arte Automobili — {title} {num}</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>{font_css(FONT_REL)}{BASE_CSS}
@page{{size:A4;margin:0}}
body{{background:#cfccc4}}
.note{{max-width:210mm;margin:8mm auto 0;font:12px/1.5 'DM Mono',monospace;color:#262B29}}
.page{{width:210mm;min-height:297mm;margin:6mm auto 12mm;background:#fff;padding:18mm 20mm 16mm;position:relative;display:flex;flex-direction:column}}
header{{display:flex;justify-content:space-between;align-items:flex-end;padding-bottom:6mm;border-bottom:.6pt solid var(--nero)}}
header .doc{{text-align:right}}
header .doc h1{{margin:0 0 1.5mm;font:400 22pt/1 'Bodoni Moda',serif}}
header .doc p{{margin:0;font:400 8.5pt/1.5 'DM Mono',monospace}}
.parties{{display:grid;grid-template-columns:1fr 1fr;gap:10mm;margin:8mm 0 6mm}}
.parties p{{margin:0;font-size:9.5pt;line-height:1.5}}
.parties .eyebrow{{margin-bottom:2mm;color:#555}}
.vehicule{{display:flex;gap:8mm;padding:3.5mm 0;border-top:.4pt solid var(--travertino);border-bottom:.4pt solid var(--travertino);font-size:9pt}}
.vehicule span.mono{{font-size:8pt;color:#555;text-transform:uppercase}}
table{{width:100%;border-collapse:collapse;margin:7mm 0 4mm;font-size:9.5pt}}
th{{text-align:left;font:600 6.5pt/1 Archivo,sans-serif;font-stretch:125%;letter-spacing:.2em;text-transform:uppercase;padding:0 0 3mm;border-bottom:.6pt solid var(--nero)}}
td{{padding:3.2mm 0;border-bottom:.4pt solid var(--travertino);vertical-align:top}}
.num{{text-align:right;font-family:'DM Mono',monospace;font-size:8.8pt;white-space:nowrap;padding-left:4mm}}
th.num{{font-family:Archivo,sans-serif}}
.totaux{{margin-left:auto;width:78mm;font-size:9.5pt}}
.totaux div{{display:flex;justify-content:space-between;padding:1.6mm 0}}
.totaux .ttc{{border-top:.6pt solid var(--nero);margin-top:1mm;padding-top:3mm;font:400 15pt/1 'Bodoni Moda',serif}}
.totaux .ttc span:last-child{{font:400 12pt/1 'DM Mono',monospace}}
.cond{{display:grid;grid-template-columns:repeat(3,1fr);gap:6mm;margin-top:9mm}}
.cond p{{margin:0;font-size:9pt;line-height:1.45}}
.cond .eyebrow{{margin-bottom:1.6mm;color:#555}}
.accord{{margin-top:8mm;border:.4pt solid var(--travertino);padding:4mm 5mm;font-size:8.5pt}}
.accord p{{margin:0 0 1mm}}.accord .sig{{height:16mm}}
footer{{margin-top:auto;padding-top:5mm;border-top:.6pt solid var(--nero);display:flex;justify-content:space-between;gap:8mm;font:400 6.8pt/1.55 'DM Mono',monospace;color:#444}}
[contenteditable]:hover{{outline:1px dashed #8E9592;outline-offset:2px}}
@media print{{body{{background:none}}.note{{display:none}}.page{{margin:0;min-height:297mm;height:297mm}}[contenteditable]:hover{{outline:none}}}}
</style></head><body>
<p class="note">Modèle modifiable : cliquez sur un texte pour le changer (quantités et prix recalculent les totaux), puis Fichier → Imprimer → Enregistrer en PDF, format A4, marges « aucune ».</p>
<main class="page">
  <header>
    {logo_svg("logo-principal", N, height=44)}
    <div class="doc"><h1>{title}</h1><p contenteditable="true">N° {num}</p><p contenteditable="true">Émis le 24.10.2026</p></div>
  </header>
  <section class="parties">
    <div><p class="eyebrow">Atelier</p><p contenteditable="true">Arte Automobili<br>[Adresse de l'atelier]<br>[Code postal, ville]<br>SIRET [—] · TVA [—]</p></div>
    <div><p class="eyebrow">Client</p><p contenteditable="true">M. Julien Martin<br>12, rue des Ateliers<br>75000 Ville<br>julien.martin@exemple.fr</p></div>
  </section>
  <section class="vehicule">
    <div><span class="mono">Véhicule</span><br><span contenteditable="true">Coupé GT · Bianco perla</span></div>
    <div><span class="mono">Immatriculation</span><br><span contenteditable="true">AB-123-CD</span></div>
    <div><span class="mono">Référence</span><br><span contenteditable="true">Opera N° 0147</span></div>
  </section>
  <table><thead><tr><th>Désignation</th><th class="num">Qté</th><th class="num">PU HT</th><th class="num">TVA</th><th class="num">Total HT</th></tr></thead>
  <tbody>{rows}</tbody></table>
  <section class="totaux">
    <div><span>Total HT</span><span class="mono" id="ht">{money(ht)}</span></div>
    <div><span>TVA 20 %</span><span class="mono" id="tva">{money(ht * .2)}</span></div>
    <div class="ttc"><span>Total TTC</span><span id="ttc">{money(ht * 1.2)}</span></div>
  </section>
  {extra_q}
  <footer>
    <div contenteditable="true">Arte Automobili · [forme juridique] au capital de [—] € · RCS [—] · SIRET [—] · TVA intracommunautaire [—]<br>
    Pénalités de retard : 3 fois le taux d'intérêt légal. Indemnité forfaitaire pour frais de recouvrement (clients professionnels) : 40 €.<br>
    Mentions à faire valider par votre expert-comptable.</div>
    {logo_svg("monogramme-micro", N, height=22)}
  </footer>
</main>
<script>
const p=s=>parseFloat(String(s).replace(/[^0-9,.-]/g,'').replace(/\\s/g,'').replace(',', '.'))||0;
const f=v=>v.toLocaleString('fr-FR',{{minimumFractionDigits:2,maximumFractionDigits:2}})+' €';
function recalc(){{let ht=0;document.querySelectorAll('tbody tr').forEach(r=>{{const t=p(r.querySelector('.qty').textContent)*p(r.querySelector('.pu').textContent);r.querySelector('.tot').textContent=f(t);ht+=t;}});
document.getElementById('ht').textContent=f(ht);document.getElementById('tva').textContent=f(ht*.2);document.getElementById('ttc').textContent=f(ht*1.2);}}
document.querySelector('tbody').addEventListener('input',recalc);
</script>
</body></html>"""


# ---------------------------------------------------------------- signature e-mail
def email_signature():
    return f"""<!doctype html>
<html lang="fr"><head><meta charset="utf-8"><title>Signature e-mail — Arte Automobili</title></head>
<body style="margin:24px;background:#ffffff">
<!-- Copier le tableau ci-dessous dans la signature de votre messagerie.
     Héberger l'image du logo sur votre site (https://arteautomobili.fr/img/signature-logo.png) et remplacer src. -->
<table cellpadding="0" cellspacing="0" border="0" role="presentation" style="font-family:Archivo,Arial,Helvetica,sans-serif;color:#0D0F0E;border-collapse:collapse">
  <tr><td style="padding:0 0 10px 0">
    <span style="font-family:'Bodoni Moda',Didot,Georgia,serif;font-size:18px;line-height:1.2">Prénom Nom</span><br>
    <span style="font-size:10px;letter-spacing:2.4px;text-transform:uppercase;color:#555">Fondateur · Arte Automobili</span>
  </td></tr>
  <tr><td style="border-top:1px solid #0D0F0E;padding:10px 0 0 0;font-family:'DM Mono',Consolas,monospace;font-size:12px;line-height:1.7;color:#262B29">
    +33 6 00 00 00 00<br>
    <a href="https://arteautomobili.fr" style="color:#1F3D35;text-decoration:none">arteautomobili.fr</a> · Sur rendez-vous
  </td></tr>
  <tr><td style="padding:14px 0 0 0">
    <img src="../../02_logo/png/aa-logo-principal-nero-1200px.png" width="150" alt="Arte Automobili" style="display:block;border:0;width:150px;height:auto">
  </td></tr>
  <tr><td style="padding:10px 0 0 0;font-family:'Bodoni Moda',Didot,Georgia,serif;font-style:italic;font-size:13px;color:#555">{SIGNATURE}</td></tr>
</table>
</body></html>"""


# ---------------------------------------------------------------- carte des prestations (A4, 2 pages)
FORMULES = [
    ("Essenziale", "dès 249 €", ["Lavage d'entretien complet", "Décontamination", "Protection polymère 6 mois"]),
    ("Completo", "dès 990 €", ["Décontamination", "Correction de vernis 1 étape", "Céramique 2 ans", "Jantes et vitrage"]),
    ("Opera", "dès 1 890 €", ["Décontamination", "Correction 2 à 3 étapes", "Céramique 5 ans", "Jantes, étriers, vitrage",
                             "Habitacle et cuir", "Fiche Opera et suivi annuel"]),
]


def brochure():
    svc = "".join(f"""<article class="svc"><div class="svc-h">{icon(p['icone'], N, 30)}<span class="mono n">{p['n']}</span></div>
      <h3>{p['nom']}</h3><p>{p['desc']}</p><p class="mono data"><span>{p['duree']}</span><span>{p['prix']}</span></p></article>"""
                  for p in PRESTATIONS)
    meth = "".join(f'<li><span class="mono">{i + 1:02d}</span><b>{t}</b><p>{d}</p></li>' for i, (t, d) in enumerate(METHODE))
    form = "".join(f"""<div class="f"><p class="eyebrow">Formule</p><h3>{n}</h3><p class="mono prix">{px}</p>
      <ul>{''.join(f'<li>{x}</li>' for x in items)}</ul></div>""" for n, px, items in FORMULES)
    return f"""<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>Carte des prestations — Arte Automobili</title>
<style>{font_css(FONT_REL)}{BASE_CSS}
@page{{size:A4;margin:0}}
.page{{width:210mm;height:297mm;overflow:hidden;position:relative;break-after:page}}
.p1{{background:var(--nero);color:var(--bianco)}}
.p1 .img{{position:absolute;inset:0 0 auto 0;height:183.5mm;background:url('{IMG_REL}rendu-07-capot-lames-de-lumiere-1080x1350.jpg') center 70%/cover}}
.p1 .lin{{position:absolute;top:183.5mm;left:0;right:0;border-top:.6pt solid var(--argento)}}
.p1 .logo{{position:absolute;top:16mm;left:18mm}}
.p1 .txt{{position:absolute;top:196mm;left:18mm;right:18mm}}
.p1 h1{{font:400 30pt/1.05 'Bodoni Moda',serif;margin:0 0 5mm}} .p1 h1 i{{font-style:italic}}
.p1 .lead{{font-size:10pt;line-height:1.6;max-width:120mm;opacity:.85;margin:0 0 8mm}}
.p1 ol{{list-style:none;margin:0;padding:0;display:grid;grid-template-columns:repeat(4,1fr);gap:5mm}}
.p1 ol li span{{display:block;color:var(--argento);font-size:8pt;margin-bottom:1.5mm}} .p1 ol li b{{font-weight:500;font-size:9.5pt}}
.p1 ol li p{{font-size:7.6pt;line-height:1.45;opacity:.75;margin:1mm 0 0}}
.p2{{background:#fff;padding:16mm 18mm}}
.p2 header{{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:.6pt solid var(--nero);padding-bottom:5mm}}
.p2 h2{{font:400 22pt/1 'Bodoni Moda',serif;margin:0}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:6mm 10mm;margin-top:8mm}}
.svc{{border-bottom:.4pt solid var(--travertino);padding-bottom:4.5mm}}
.svc-h{{display:flex;justify-content:space-between;align-items:center}} .svc .n{{font-size:8pt;color:#666}}
.svc h3{{font:400 13pt/1.2 'Bodoni Moda',serif;margin:2.5mm 0 1.5mm}}
.svc p{{font-size:8.2pt;line-height:1.5;margin:0}}
.svc .data{{display:flex;justify-content:space-between;margin-top:2.5mm;font-size:8pt;text-transform:uppercase}}
.form{{display:grid;grid-template-columns:repeat(3,1fr);gap:5mm;margin-top:8mm}}
.f{{background:var(--nero);color:var(--bianco);padding:5mm}} .f:nth-child(3){{background:var(--verde)}}
.f .eyebrow{{color:var(--argento);margin:0}} .f h3{{font:italic 400 16pt/1.1 'Bodoni Moda',serif;margin:2mm 0 1mm}}
.f .prix{{font-size:9pt;margin:0 0 3mm}} .f ul{{margin:0;padding:0;list-style:none;font-size:7.8pt;line-height:1.7}}
.f li{{border-top:.4pt solid rgba(238,237,231,.18)}}
.note{{font:7pt/1.5 'DM Mono',monospace;color:#666;margin-top:5mm}}
</style></head><body>
<section class="page p1"><div class="img"></div><div class="lin"></div>
 <div class="logo">{logo_svg("logo-principal", B, height=46)}</div>
 <div class="txt"><p class="eyebrow" style="color:{A};margin:0 0 4mm">Carte des prestations · 2026</p>
 <h1>L'arte del <i>riflesso.</i></h1><p class="lead">{PITCH_COURT}</p><ol>{meth}</ol></div></section>
<section class="page p2"><header><h2>Prestations</h2><span class="mono" style="font-size:8pt">TARIFS INDICATIFS TTC · DEVIS APRÈS INSPECTION</span></header>
 <div class="grid">{svc}</div><div class="form">{form}</div>
 <p class="note">Les tarifs dépendent du gabarit et de l'état du véhicule. Durées données à titre indicatif.</p></section>
</body></html>"""


# ---------------------------------------------------------------- présentation commerciale 16:9
def deck():
    S = "section"
    svc = "".join(f'<div class="s">{icon(p["icone"], B, 40)}<h3>{p["nom"]}</h3><p class="mono">{p["duree"]} · {p["prix"]}</p></div>'
                  for p in PRESTATIONS)
    meth = "".join(f'<div class="m"><span class="mono">{i + 1:02d}</span><h3>{t}</h3><p>{d}</p></div>' for i, (t, d) in enumerate(METHODE))
    avis = "".join(f'<blockquote><p>« {q} »</p><cite>{n}<span class="mono">{c}</span></cite></blockquote>' for n, c, q in AVIS)
    return f"""<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>Présentation — Arte Automobili</title>
<style>{font_css(FONT_REL)}{BASE_CSS}
@page{{size:1920px 1080px;margin:0}}
{S}{{width:1920px;height:1080px;position:relative;overflow:hidden;background:var(--nero);color:var(--bianco);break-after:page}}
.lin{{position:absolute;left:0;right:0;top:667px;border-top:1.5px solid var(--argento)}}
.ey{{position:absolute;left:120px;top:96px;color:var(--argento);font-size:16px}}
.pg{{position:absolute;right:120px;top:96px;font:400 16px 'DM Mono',monospace;color:var(--argento)}}
h1,h2{{font-family:'Bodoni Moda',serif;font-weight:400;margin:0}} h2{{font-size:76px;line-height:1.05}}
i{{font-style:italic}}
.cover{{background:url('{IMG_REL}rendu-01-hero-studio-2400x1200.jpg') right center/cover}}
.cover::before{{content:"";position:absolute;inset:0;background:linear-gradient(90deg,rgba(13,15,14,.95) 25%,rgba(13,15,14,0) 65%)}}
.cover .in{{position:absolute;left:120px;top:150px}} .cover h1{{font-size:96px;line-height:1.02;margin-top:300px}}
.cover p{{font:400 20px 'DM Mono',monospace;color:var(--argento);margin-top:40px;letter-spacing:.06em}}
.split .img{{position:absolute;left:0;top:0;width:820px;height:1080px;background:center/cover}}
.split .tx{{position:absolute;left:960px;right:120px;top:200px}}
.split p.l{{font-size:24px;line-height:1.6;opacity:.85;margin:36px 0 0;max-width:760px}}
.ms{{position:absolute;left:120px;right:120px;top:720px;display:grid;grid-template-columns:repeat(4,1fr);gap:48px}}
.m span{{font-size:18px;color:var(--argento)}} .m h3{{font:400 34px 'Bodoni Moda',serif;margin:14px 0 10px}} .m p{{font-size:19px;line-height:1.5;opacity:.75;margin:0}}
.ttl{{position:absolute;left:120px;top:420px}}
.svcs{{position:absolute;left:120px;right:120px;top:720px;display:grid;grid-template-columns:repeat(6,1fr);gap:32px}}
.s h3{{font:400 26px/1.15 'Bodoni Moda',serif;margin:18px 0 10px}} .s p{{font-size:15px;color:var(--argento);margin:0;text-transform:uppercase}}
.ba .a,.ba .b{{position:absolute;top:0;height:1080px;width:960px;background:center 35%/cover}}
.ba .a{{left:0}} .ba .b{{left:960px}} .ba .v{{position:absolute;left:959px;top:0;bottom:0;border-left:2px solid var(--bianco)}}
.ba .lab{{position:absolute;top:96px;font-size:18px;color:var(--bianco)}}
.ba .cap{{position:absolute;left:120px;bottom:96px;font:400 18px 'DM Mono',monospace}}
.av{{position:absolute;left:120px;right:120px;top:300px;display:grid;grid-template-columns:repeat(3,1fr);gap:64px}}
blockquote{{margin:0}} blockquote p{{font:italic 400 38px/1.25 'Bodoni Moda',serif;margin:0 0 32px}}
cite{{font-style:normal;font-size:20px}} cite span{{display:block;color:var(--argento);font-size:15px;margin-top:8px;text-transform:uppercase}}
.contact{{display:flex;flex-direction:column;align-items:center;justify-content:center;text-align:center}}
.contact p{{font:400 22px/1.9 'DM Mono',monospace;color:var(--bianco);margin:0}}
.verde{{background:var(--verde)}}
.op .card{{position:absolute;left:1060px;top:200px;width:700px;height:497px;background:var(--bianco);color:var(--nero);padding:56px;box-shadow:0 40px 80px rgba(0,0,0,.5)}}
.op .card .t{{font:italic 400 64px 'Bodoni Moda',serif}} .op .card .t span{{font:400 34px 'DM Mono',monospace;font-style:normal;margin-left:18px}}
.op .card dl{{display:grid;grid-template-columns:200px 1fr;gap:14px 0;margin:40px 0 0;font-size:20px}} .op .card dt{{font:400 14px 'DM Mono',monospace;text-transform:uppercase;color:#666;padding-top:5px}} .op .card dd{{margin:0}}
</style></head><body>
<{S} class="cover"><div class="in">{logo_svg("logo-principal", B, height=96)}<h1>L'arte del<br><i>riflesso.</i></h1><p>PRÉSENTATION · 2026</p></div></{S}>
<{S} class="split"><div class="img" style="background-image:url('{IMG_REL}rendu-07-capot-lames-de-lumiere-1080x1350.jpg')"></div>
 <div class="tx"><p class="eyebrow" style="color:{A};font-size:16px">L'atelier</p><h2 style="margin-top:28px">Une automobile se juge<br><i>à son reflet.</i></h2><p class="l">{PITCH_LONG[0]}</p><p class="l">{PITCH_LONG[1]}</p></div><span class="pg">02</span></{S}>
<{S}><p class="eyebrow ey">La méthode</p><span class="pg">03</span><div class="ttl"><h2>Quatre temps,<br><i>aucun raccourci.</i></h2></div><div class="lin"></div><div class="ms">{meth}</div></{S}>
<{S}><p class="eyebrow ey">Prestations</p><span class="pg">04</span><div class="ttl"><h2>Six savoir-faire,<br><i>un seul protocole.</i></h2></div><div class="lin"></div><div class="svcs">{svc}</div></{S}>
<{S} class="ba"><div class="a" style="background-image:url('{IMG_REL}rendu-04-epaule-avant-1080x1350.jpg')"></div><div class="b" style="background-image:url('{IMG_REL}rendu-03-epaule-apres-1080x1350.jpg')"></div><div class="v"></div>
 <p class="eyebrow lab" style="left:120px">Avant</p><p class="eyebrow lab" style="left:1080px">Après</p><p class="cap">MÊME LUMIÈRE, MÊME ANGLE · OPERA N° 0052</p></{S}>
<{S} class="op verde"><p class="eyebrow ey" style="color:{B};opacity:.7">Le rituel</p><span class="pg" style="color:{B}">06</span>
 <div style="position:absolute;left:120px;top:300px;width:820px"><h2>Chaque voiture reçoit<br><i>son numéro d'Opera.</i></h2><p style="font-size:24px;line-height:1.6;opacity:.85;margin-top:36px">À la restitution, une fiche détaille les interventions, les produits utilisés et le calendrier d'entretien conseillé. Le numéro suit la voiture d'une visite à l'autre.</p></div>
 <div class="card"><div class="t">Opera<span>N° 0147</span></div><dl><dt>Restitué le</dt><dd>24.10.2026</dd><dt>Prestations</dt><dd>Correction 2 étapes · Céramique 5 ans</dd><dt>Prochain contrôle</dt><dd>Avril 2027</dd></dl></div></{S}>
<{S}><p class="eyebrow ey">Ils en parlent</p><span class="pg">07</span><div class="av">{avis}</div><div class="lin"></div></{S}>
<{S} class="contact">{logo_svg("logo-vertical", B, height=300)}<div style="height:72px"></div><p>SUR RENDEZ-VOUS · [ADRESSE DE L'ATELIER]</p><p>+33 6 00 00 00 00 · CONTACT@ARTEAUTOMOBILI.FR</p></{S}>
</body></html>"""


# ---------------------------------------------------------------- en-tête Word
def word_letterhead(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(48), Mm(30)
    sec.header_distance, sec.footer_distance = Mm(16), Mm(12)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Archivo", Pt(10)
    st.font.color.rgb = RGBColor(0x0D, 0x0F, 0x0E)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.35
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(str(ROOT / "02_logo" / "png" / "aa-logo-principal-nero-1200px.png"), height=Mm(11.5))
    fp = sec.footer.paragraphs[0]
    r = fp.add_run("Arte Automobili · [forme juridique] au capital de [—] € · SIRET [—] · TVA [—]\n"
                   "[Adresse de l'atelier] · +33 6 00 00 00 00 · contact@arteautomobili.fr · arteautomobili.fr")
    r.font.name, r.font.size = "DM Mono", Pt(6.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = d.add_paragraph("Ville, le [date]")
    p.runs[0].font.name = "DM Mono"
    d.add_paragraph()
    t = d.add_paragraph("Objet : [objet du courrier]")
    t.runs[0].bold = True
    d.add_paragraph("Madame, Monsieur,")
    d.add_paragraph("[Corps du courrier — Archivo 10 pt, interlignage 1,35. Polices à installer depuis "
                    "03_typographie/polices ; à défaut Word utilisera Arial.]")
    s = d.add_paragraph("Prénom Nom")
    s.runs[0].font.name, s.runs[0].font.size = "Bodoni Moda", Pt(13)
    d.save(path)


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    pages = {"devis-modele.html": invoice("devis"), "facture-modele.html": invoice("facture"),
             "carte-des-prestations.html": brochure(), "presentation-commerciale.html": deck()}
    for k, v in pages.items():
        (DOCS / k).write_text(v)
    (DOCS / "signature-email.html").write_text(email_signature())
    word_letterhead(DOCS / "papier-en-tete.docx")
    jobs = [
        {"in": str(DOCS / "devis-modele.html"), "out": str(DOCS / "devis-exemple.pdf"), "w": 900, "h": 1200, "pdf": {"width": "210mm", "height": "297mm"}},
        {"in": str(DOCS / "facture-modele.html"), "out": str(DOCS / "facture-exemple.pdf"), "w": 900, "h": 1200, "pdf": {"width": "210mm", "height": "297mm"}},
        {"in": str(DOCS / "carte-des-prestations.html"), "out": str(DOCS / "carte-des-prestations-a4.pdf"), "w": 900, "h": 1200, "pdf": {"width": "210mm", "height": "297mm"}},
        {"in": str(DOCS / "presentation-commerciale.html"), "out": str(DOCS / "presentation-commerciale-16x9.pdf"), "w": 1920, "h": 1080, "pdf": {"width": "1920px", "height": "1080px"}},
        {"in": str(DOCS / "signature-email.html"), "out": str(DOCS / "signature-email.png"), "w": 520, "h": 330, "scale": 2},
    ]
    jf = DOCS / "_jobs.json"
    jf.write_text(json.dumps(jobs))
    subprocess.run(["node", str(ROOT / "_source" / "render.mjs"), str(jf)], check=True, stdout=subprocess.DEVNULL)
    jf.unlink()
    print("documents ok")


if __name__ == "__main__":
    main()
